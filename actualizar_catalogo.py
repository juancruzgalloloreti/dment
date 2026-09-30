"""
D'MENT — Conversor de catalogo a JSON
======================================
Uso:
  python actualizar_catalogo.py

Lee el archivo "catalogo_productos.csv" (o .xlsx si tenes openpyxl instalado)
y genera "productos.json" que la pagina web usa automaticamente.

COLUMNAS DEL CSV/EXCEL:
  id          | numero unico (ej: 1, 2, 3...)
  nombre      | nombre del producto (ej: Remera Basica)
  categoria   | Remeras / Musculosas / Buzos / Camperas / Pantalones / Shorts
  descripcion | descripcion corta
  precio      | precio mayorista en pesos (solo numero, ej: 3500)
  imagen      | URL de la imagen o ruta local (puede quedar vacio)
  destacado   | SI o NO
  colores     | separados por | (ej: Negro|Blanco|Gris)
  talles      | separados por | (ej: S|M|L|XL)
  codigo      | codigo interno (ej: REM-001)
"""

import csv
import json
import os
import sys

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
CSV_PATH   = os.path.join(SCRIPT_DIR, 'catalogo_productos.csv')
XLSX_PATH  = os.path.join(SCRIPT_DIR, 'catalogo_productos.xlsx')
JSON_PATH  = os.path.join(SCRIPT_DIR, 'productos.json')


def parse_bool(val):
    return str(val).strip().upper() in ('SI', 'YES', 'TRUE', '1', 'S')

def parse_list(val):
    if not val:
        return []
    return [v.strip() for v in str(val).split('|') if v.strip()]

def load_csv(path):
    rows = []
    with open(path, encoding='utf-8-sig', newline='') as f:
        reader = csv.DictReader(f)
        for row in reader:
            rows.append(dict(row))
    return rows

def load_xlsx(path):
    try:
        import openpyxl
    except ImportError:
        print("Para leer .xlsx instala openpyxl:  pip install openpyxl")
        sys.exit(1)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = {headers[i]: (row[i] if row[i] is not None else '') for i in range(len(headers))}
        rows.append(d)
    return rows

def convert(rows):
    products = []
    for r in rows:
        nombre = str(r.get('nombre', '')).strip()
        if not nombre:
            continue
        try:
            precio = float(str(r.get('precio', '0')).replace(',', '.'))
        except:
            precio = 0
        products.append({
            'id':          r.get('id', ''),
            'nombre':      nombre,
            'categoria':   str(r.get('categoria', '')).strip(),
            'descripcion': str(r.get('descripcion', '')).strip(),
            'precio':      precio,
            'imagen':      str(r.get('imagen', '')).strip(),
            'destacado':   parse_bool(r.get('destacado', 'NO')),
            'colores':     parse_list(r.get('colores', '')),
            'talles':      parse_list(r.get('talles', '')),
            'codigo':      str(r.get('codigo', '')).strip(),
        })
    return products

def main():
    # Prioridad: xlsx > csv
    if os.path.exists(XLSX_PATH):
        print(f"Leyendo: {XLSX_PATH}")
        rows = load_xlsx(XLSX_PATH)
    elif os.path.exists(CSV_PATH):
        print(f"Leyendo: {CSV_PATH}")
        rows = load_csv(CSV_PATH)
    else:
        print("ERROR: No se encontro 'catalogo_productos.csv' ni 'catalogo_productos.xlsx'")
        sys.exit(1)

    products = convert(rows)
    with open(JSON_PATH, 'w', encoding='utf-8') as f:
        json.dump(products, f, ensure_ascii=False, indent=2)

    print(f"OK: Se exportaron {len(products)} productos a 'productos.json'")
    print("\nProductos:")
    for p in products:
        print(f"  [{p['categoria']}] {p['nombre']} - ${p['precio']:.0f}")

if __name__ == '__main__':
    main()
